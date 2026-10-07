"""Generate and publish the daily pilot workbook using raw weighted weekday demand.

This publisher is intentionally read-only with respect to ClickHouse forecast
runs. It uses the current bakery/SKU assortment table as the publishable scope,
then sets the SKU forecast quantity to a weighted average of restored demand
from previous same-weekday observations.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
import time as _time
import urllib.request
from datetime import date as date_type
from datetime import timedelta
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "apps"))
sys.path.insert(0, str(ROOT / "apps" / "forecast_embedded"))

from app.db import get_client  # noqa: E402
from app.table_names import table_name  # noqa: E402
from scripts.purchased_product_metadata import (  # noqa: E402
    MISSING_KRATNOST_LABEL,
    MISSING_SHELF_LIFE_LABEL,
    PURCHASED_CATEGORIES,
    get_purchased_product_metadata,
)


WEIGHTS = [1.0, 1.15, 1.35, 1.65, 2.0]
PUBLISHABLE_PATTERN = "пирог|выпеч|фастфуд|хлеб|пирожн|маффин|печенье|донат|торт|рулет"
WEEKDAY_RU = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
BAKEABLE_CATEGORIES = {"Пироги сытные", "Пироги сладкие", "Выпечка сытная", "Выпечка сладкая", "Фастфуд"}
PUBLISHABLE_CATEGORIES = BAKEABLE_CATEGORIES | PURCHASED_CATEGORIES
PRODUCT_NAME_OVERRIDES = {
    11615: "Плетенка кленовая",
    11616: "Плетенка с черникой",
    11617: "Плетенка с земляникой",
}
MISSING_STOCK_LABEL = "нет данных по остатку"
_PILOT_SCOPE_NAME = "expanded_pilot_38"
_SEED_PILOT_IDS = [
    1, 3, 12, 13, 14, 20, 21, 22, 26, 27, 28, 39, 41, 56, 57, 66, 67, 69,
    80, 89, 99, 107, 113, 125, 149, 153, 155, 160, 191, 221, 222, 229, 230,
    246, 257, 260, 268, 270,
]

VIBECODE_API_BASE = "https://vibecode.bitrix24.tech/v1"
PILOT_CHAT_DIALOG_ID = "chat179919"
PILOT_CHAT_ID = 179919
PILOT_CHAT_DISK_FOLDER_ID = 1473995
B24_WEBHOOK_URL_ENV = "B24_WEBHOOK_URL"


def load_env(path: str | Path) -> None:
    env_path = Path(path)
    if not env_path.exists():
        return
    for raw in env_path.read_text(encoding="utf-8-sig").splitlines():
        if "=" in raw and not raw.lstrip().startswith("#"):
            key, value = raw.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def env_flag(name: str, *, default: bool = False) -> bool:
    value = os.environ.get(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def weighted_tail(values: pd.Series) -> float:
    tail = values.tail(len(WEIGHTS)).astype(float)
    if tail.empty:
        return 0.0
    weights = np.asarray(WEIGHTS[-len(tail) :], dtype=float)
    return float((tail.to_numpy(dtype=float) * weights).sum() / weights.sum())


def load_pilot_bakery_ids(client) -> list[int]:
    try:
        df = client.query_df(
            """
            SELECT bakery_id, argMax(action, changed_at) AS last_action
            FROM Svezhar.pilot_scope_events
            WHERE scope_name = %(scope_name)s
            GROUP BY bakery_id
            """,
            parameters={"scope_name": _PILOT_SCOPE_NAME},
        )
    except Exception as exc:
        print(f"  WARNING: cannot load pilot_scope_events, using seed: {exc}")
        return list(_SEED_PILOT_IDS)
    if df.empty:
        print("  INFO: pilot_scope_events empty, using seed")
        return list(_SEED_PILOT_IDS)
    event_all = {int(row["bakery_id"]) for row in df.to_dict("records")}
    event_active = {
        int(row["bakery_id"])
        for row in df.to_dict("records")
        if str(row.get("last_action") or "").lower() not in {"exclude", "remove", "delete"}
    }
    seed_active = {bid for bid in _SEED_PILOT_IDS if bid not in event_all}
    result = sorted(event_active | seed_active)
    print(f"  INFO: pilot — {len(result)} bakeries ({len(event_active)} from events, {len(seed_active)} from seed)")
    return result


def fetch_active_scope(client, forecast_date: str, pilot_bakery_ids: list[int]) -> tuple[str, pd.DataFrame]:
    run_df = client.query_df(
        f"""
        select run_id, model_version
        from {table_name('forecast_runs_embedded')}
        where status = 'active'
        limit 1
        """
    )
    if run_df.empty:
        raise RuntimeError("No active forecast run")
    run_id = str(run_df.iloc[0]["run_id"])
    scope = client.query_df(
        """
        with latest as (
            select toInt64(bakery_id) as bakery_id, max(valid_from) as latest_valid_from
            from Svezhar.bakery_product_assortment_embedded final
            where valid_from <= toDate(%(forecast_date)s)
              and toInt64(bakery_id) in %(bids)s
            group by bakery_id
        ),
        products as (
            select
                toInt64OrZero(toString(product_id)) as product_id,
                argMax(product_name, _updated_at) as product_name,
                argMax(category_name, _updated_at) as category_name
            from Svezhar.dim_products
            group by product_id
        )
        select
            toInt64(a.bakery_id) as bakery_id,
            toInt64OrZero(toString(a.product_id)) as product_id,
            any(p.product_name) as product_name,
            any(p.category_name) as category_name,
            0.0 as active_forecast_qty
        from Svezhar.bakery_product_assortment_embedded as a final
        inner join latest l
          on toInt64(a.bakery_id) = l.bakery_id
         and a.valid_from = l.latest_valid_from
        inner join products p
          on p.product_id = toInt64OrZero(toString(a.product_id))
        group by bakery_id, product_id
        """,
        parameters={"forecast_date": forecast_date, "bids": pilot_bakery_ids},
    )
    if scope.empty:
        raise RuntimeError(f"No publishable assortment rows for {forecast_date}, active_run={run_id}")
    scope["category_name"] = scope["category_name"].fillna("")
    scope = scope[scope["category_name"].str.lower().str.contains(PUBLISHABLE_PATTERN, regex=True, na=False)].copy()
    if scope.empty:
        raise RuntimeError(f"No publishable assortment rows after category filter for {forecast_date}, active_run={run_id}")
    return f"assortment_table; active_run={run_id}", scope


def fetch_history_demand(client, history_dates: list[str], pilot_bakery_ids: list[int]) -> pd.DataFrame:
    params = {"dates": history_dates, "bids": pilot_bakery_ids}
    sales = client.query_df(
        """
        select check_date date, toInt64OrZero(toString(bakery_id)) bakery_id,
               toInt64OrZero(toString(product_id)) product_id,
               sum(toFloat64(quantity)) observed_sales_qty,
               min(check_datetime) first_sale_time,
               max(check_datetime) last_sale_time
        from (
            select distinct check_datetime, check_date, bakery_id, product_id, quantity, line_amount
            from Svezhar.fct_check_lines
            where hex(cash_event_type)='D09FD180D0BED0B4D0B0D0B6D0B0'
              and check_date in %(dates)s
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    release = client.query_df(
        """
        select rd date, toInt64OrZero(toString(bid)) bakery_id,
               toInt64OrZero(toString(pid)) product_id, sum(qty) release_qty
        from (
            select argMax(release_date,_updated_at) rd, argMax(bakery_id,_updated_at) bid,
                   argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                   argMax(is_deleted,_updated_at) deleted
            from Svezhar.fct_production_release
            where release_date in %(dates)s
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
            group by release_id,line_id
            having deleted not in ('1','true','Да')
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    moves = client.query_df(
        """
        select date, bakery_id, product_id,
               sum(incoming_move_qty) incoming_move_qty,
               sum(outgoing_move_qty) outgoing_move_qty
        from (
            select md date, toInt64OrZero(toString(receiver)) bakery_id,
                   toInt64OrZero(toString(pid)) product_id, qty incoming_move_qty, 0.0 outgoing_move_qty
            from (
                select argMax(move_date,_updated_at) md, argMax(receiver_id,_updated_at) receiver,
                       argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                       argMax(is_deleted,_updated_at) deleted
                from Svezhar.fct_moves
                where move_date in %(dates)s
                group by move_id,line_id
                having deleted not in ('1','true','Да')
            )
            union all
            select md date, toInt64OrZero(toString(sender)) bakery_id,
                   toInt64OrZero(toString(pid)) product_id, 0.0 incoming_move_qty, qty outgoing_move_qty
            from (
                select argMax(move_date,_updated_at) md, argMax(sender_id,_updated_at) sender,
                       argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                       argMax(is_deleted,_updated_at) deleted
                from Svezhar.fct_moves
                where move_date in %(dates)s
                group by move_id,line_id
                having deleted not in ('1','true','Да')
            )
        )
        where bakery_id in %(bids)s
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    writeoffs = client.query_df(
        """
        select wd date, toInt64OrZero(toString(bid)) bakery_id,
               toInt64OrZero(toString(pid)) product_id, sum(qty) written_off_qty
        from (
            select argMax(write_off_date,_updated_at) wd, argMax(bakery_id,_updated_at) bid,
                   argMax(write_off_product_id,_updated_at) pid,
                   toFloat64(argMax(write_off_qty,_updated_at)) qty,
                   argMax(is_deleted,_updated_at) deleted
            from Svezhar.fct_write_offs
            where write_off_date in %(dates)s
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
            group by write_off_doc_num,line_id
            having deleted not in ('1','true','Да')
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    keys = ["date", "bakery_id", "product_id"]
    facts = sales
    for part in (release, moves, writeoffs):
        facts = facts.merge(part, on=keys, how="outer")
    facts["date"] = pd.to_datetime(facts["date"]).dt.normalize()
    for column in ["observed_sales_qty", "release_qty", "incoming_move_qty", "outgoing_move_qty", "written_off_qty"]:
        facts[column] = pd.to_numeric(facts.get(column, 0.0), errors="coerce").fillna(0.0).clip(lower=0.0)
    facts["available_proxy"] = (facts["release_qty"] + facts["incoming_move_qty"] - facts["outgoing_move_qty"]).clip(lower=0.0)
    bakery_last = (
        facts.dropna(subset=["last_sale_time"])
        .groupby(["date", "bakery_id"], as_index=False)["last_sale_time"]
        .max()
        .rename(columns={"last_sale_time": "bakery_last_sale_time"})
    )
    facts = facts.merge(bakery_last, on=["date", "bakery_id"], how="left", validate="many_to_one")

    last_sale = pd.to_datetime(facts["last_sale_time"], errors="coerce", utc=True)
    bakery_end = pd.to_datetime(facts["bakery_last_sale_time"], errors="coerce", utc=True)
    opening = facts["date"].dt.tz_localize("Europe/Moscow").dt.tz_convert("UTC") + pd.Timedelta(hours=7.5)
    elapsed_hours = (last_sale - opening).dt.total_seconds() / 3600
    remaining_hours = (bakery_end - last_sale).dt.total_seconds().clip(lower=0) / 3600
    stockout_like = (
        facts["observed_sales_qty"].gt(0.0)
        & facts["available_proxy"].gt(0.0)
        & (facts["observed_sales_qty"] + facts["written_off_qty"]).ge(facts["available_proxy"] * 0.95)
        & elapsed_hours.ge(2.0)
        & remaining_hours.gt(0.0)
    )
    raw_lost = facts["observed_sales_qty"] / elapsed_hours.replace(0.0, np.nan) * remaining_hours
    cap = pd.concat([facts["observed_sales_qty"] * 1.5, pd.Series(15.0, index=facts.index)], axis=1).max(axis=1)
    facts["lost_demand_estimate"] = raw_lost.where(stockout_like, 0.0).fillna(0.0).clip(lower=0.0).clip(upper=cap)
    facts["demand"] = facts["observed_sales_qty"] + facts["lost_demand_estimate"]
    return facts[keys + ["demand"]].copy()


def build_weighted_weekday(scope: pd.DataFrame, history: pd.DataFrame, forecast_date: str) -> pd.DataFrame:
    date = pd.Timestamp(forecast_date).normalize()
    hist = history[history["date"].dt.dayofweek.eq(date.dayofweek)].sort_values("date").copy()
    current = scope[["bakery_id", "product_id", "product_name", "category_name"]].drop_duplicates().copy()
    if hist.empty:
        current["forecast_qty"] = 0.0
        return current
    pair_plan = (
        hist.groupby(["bakery_id", "product_id"], as_index=False)["demand"]
        .agg(forecast_qty=weighted_tail)
    )
    current = current.merge(pair_plan, on=["bakery_id", "product_id"], how="left")
    current["forecast_qty"] = current["forecast_qty"].fillna(0.0).clip(lower=0.0)
    return current[["bakery_id", "product_id", "product_name", "category_name", "forecast_qty"]]


def fetch_bakery_info(client, pilot_bakery_ids: list[int]) -> dict[int, dict[str, str]]:
    pilot_bids_str = [f"{b:09d}" for b in pilot_bakery_ids]
    bakery_df = client.query_df(
        """
        select bakery_id as bid, any(bakery_name) as name, any(city) as city
        from dim_bakeries
        where bakery_id in %(bids)s
        group by bakery_id
        """,
        parameters={"bids": pilot_bids_str},
    )
    result = {}
    for row in bakery_df.to_dict("records"):
        try:
            result[int(row["bid"])] = {"name": str(row["name"]), "city": str(row["city"])}
        except (TypeError, ValueError):
            continue
    return result


def round_up_kratnost(value: float, kratnost: int) -> int:
    if value <= 0 or kratnost <= 0:
        return 0
    return int(math.ceil(value / kratnost - 1e-9) * kratnost)


def plan_with_optional_kratnost(net_need: float, kratnost: int | None) -> tuple[int, int | str]:
    if kratnost is None:
        return max(0, int(math.ceil(net_need - 1e-9))), MISSING_KRATNOST_LABEL
    return round_up_kratnost(net_need, kratnost), kratnost


def fetch_stock_and_meta(
    client,
    forecast_date: str,
    forecast: pd.DataFrame,
    pilot_bakery_ids: list[int],
) -> tuple[dict[tuple[int, int], float], set[int], dict[int, int], dict[tuple[int, int], int], set[int]]:
    previous_date = str(date_type.fromisoformat(forecast_date) - timedelta(days=1))
    params = {"previous_date": previous_date, "bids": pilot_bakery_ids}
    sold = client.query_df(
        """
        select check_date date, toInt64OrZero(toString(bakery_id)) bakery_id,
               toInt64OrZero(toString(product_id)) product_id, sum(toFloat64(quantity)) qty_sold
        from (
            select distinct check_datetime, check_date, bakery_id, product_id, quantity, line_amount
            from Svezhar.fct_check_lines
            where hex(cash_event_type)='D09FD180D0BED0B4D0B0D0B6D0B0'
              and check_date = toDate(%(previous_date)s)
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    produced = client.query_df(
        """
        select rd date, toInt64OrZero(toString(bid)) bakery_id,
               toInt64OrZero(toString(pid)) product_id, sum(qty) qty_produced
        from (
            select argMax(release_date,_updated_at) rd, argMax(bakery_id,_updated_at) bid,
                   argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                   argMax(is_deleted,_updated_at) deleted
            from Svezhar.fct_production_release
            where release_date = toDate(%(previous_date)s)
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
            group by release_id,line_id
            having deleted not in ('1','true','Да')
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    moves = client.query_df(
        """
        select date, bakery_id, product_id, sum(qty_received) qty_received, sum(qty_sent) qty_sent
        from (
            select md date, toInt64OrZero(toString(receiver)) bakery_id,
                   toInt64OrZero(toString(pid)) product_id, qty qty_received, 0.0 qty_sent
            from (
                select argMax(move_date,_updated_at) md, argMax(receiver_id,_updated_at) receiver,
                       argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                       argMax(is_deleted,_updated_at) deleted
                from Svezhar.fct_moves
                where move_date = toDate(%(previous_date)s)
                group by move_id,line_id
                having deleted not in ('1','true','Да')
            )
            union all
            select md date, toInt64OrZero(toString(sender)) bakery_id,
                   toInt64OrZero(toString(pid)) product_id, 0.0 qty_received, qty qty_sent
            from (
                select argMax(move_date,_updated_at) md, argMax(sender_id,_updated_at) sender,
                       argMax(product_id,_updated_at) pid, toFloat64(argMax(quantity,_updated_at)) qty,
                       argMax(is_deleted,_updated_at) deleted
                from Svezhar.fct_moves
                where move_date = toDate(%(previous_date)s)
                group by move_id,line_id
                having deleted not in ('1','true','Да')
            )
        )
        where bakery_id in %(bids)s
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    writeoffs = client.query_df(
        """
        select wd date, toInt64OrZero(toString(bid)) bakery_id,
               toInt64OrZero(toString(pid)) product_id, sum(qty) qty_written_off
        from (
            select argMax(write_off_date,_updated_at) wd, argMax(bakery_id,_updated_at) bid,
                   argMax(write_off_product_id,_updated_at) pid,
                   toFloat64(argMax(write_off_qty,_updated_at)) qty,
                   argMax(is_deleted,_updated_at) deleted
            from Svezhar.fct_write_offs
            where write_off_date = toDate(%(previous_date)s)
              and toInt64OrZero(toString(bakery_id)) in %(bids)s
            group by write_off_doc_num,line_id
            having deleted not in ('1','true','Да')
        )
        group by date, bakery_id, product_id
        """,
        parameters=params,
    )
    keys = ["date", "bakery_id", "product_id"]
    stock = sold
    for part in (produced, moves, writeoffs):
        stock = stock.merge(part, on=keys, how="outer")
    for column in ["qty_sold", "qty_produced", "qty_received", "qty_sent", "qty_written_off"]:
        stock[column] = pd.to_numeric(stock.get(column, 0.0), errors="coerce").fillna(0.0).clip(lower=0.0)
    stock["stock_balance"] = (
        stock["qty_produced"] + stock["qty_received"] - stock["qty_sent"] - stock["qty_sold"] - stock["qty_written_off"]
    ).clip(lower=0.0)
    unavailable = set(
        stock.groupby("bakery_id", as_index=False)[["qty_sold", "qty_produced", "qty_received"]]
        .sum()
        .query("qty_sold > 0 and (qty_produced + qty_received) <= 0")["bakery_id"]
        .astype(int)
        .tolist()
    )
    stock_lookup = {
        (int(row["bakery_id"]), int(row["product_id"])): float(row["stock_balance"])
        for row in stock.to_dict("records")
    }

    product_ids = [int(pid) for pid in forecast["product_id"].dropna().unique()]
    product_ids_padded = [f"{pid:09d}" for pid in product_ids]
    meta = client.query_df(
        f"""
        select product_id, bakery_id, dough_group, kratnost, scope
        from {table_name('baking_sku_meta')} final
        where is_active = 1 and product_id in %(pids)s
        """,
        parameters={"pids": product_ids_padded},
    )
    base_kratnost: dict[int, int] = {}
    bakery_kratnost: dict[tuple[int, int], int] = {}
    frozen_pids: set[int] = set()
    for row in meta.to_dict("records"):
        try:
            pid = int(row["product_id"])
        except (TypeError, ValueError):
            continue
        if "замороженные полуфабрикаты" in str(row.get("dough_group") or "").lower():
            frozen_pids.add(pid)
            continue
        kratnost = int(row.get("kratnost") or 1) or 1
        if row.get("scope") == "bakery" and row.get("bakery_id") is not None:
            try:
                bakery_kratnost[(pid, int(row["bakery_id"]))] = kratnost
            except (TypeError, ValueError):
                pass
        else:
            base_kratnost[pid] = kratnost
    return stock_lookup, unavailable, base_kratnost, bakery_kratnost, frozen_pids


def build_rows(client, forecast_date: str, forecast: pd.DataFrame, pilot_bakery_ids: list[int]) -> list[dict]:
    bakery_info = fetch_bakery_info(client, pilot_bakery_ids)
    stock_lookup, unavailable, base_kratnost, bakery_kratnost, frozen_pids = fetch_stock_and_meta(
        client, forecast_date, forecast, pilot_bakery_ids
    )
    rows: list[dict] = []
    for rec in forecast.to_dict("records"):
        bid = int(rec["bakery_id"])
        pid = int(rec["product_id"])
        if bid not in pilot_bakery_ids:
            continue
        category = str(rec.get("category_name") or "")
        if category not in PUBLISHABLE_CATEGORIES or pid in frozen_pids:
            continue
        forecast_qty = max(float(rec.get("forecast_qty") or 0.0), 0.0)
        metadata = get_purchased_product_metadata(rec.get("product_name"))
        if category in PURCHASED_CATEGORIES:
            stock_qty = 0.0
            stock_unavailable = metadata.shelf_life is None and category != "Хлеб"
            kratnost: int | None = metadata.kratnost
        else:
            stock_qty = max(stock_lookup.get((bid, pid), 0.0), 0.0)
            stock_unavailable = bid in unavailable
            kratnost = bakery_kratnost.get((pid, bid)) or base_kratnost.get(pid)
        net_need = max(forecast_qty - stock_qty, 0.0)
        production_plan, kratnost_display = plan_with_optional_kratnost(net_need, kratnost)
        rows.append(
            {
                "bakery_id": bid,
                "bakery_name": bakery_info.get(bid, {}).get("name") or str(bid),
                "category": category,
                "product_name": PRODUCT_NAME_OVERRIDES.get(pid, str(rec.get("product_name") or "")),
                "forecast": round(forecast_qty, 1),
                "yesterday_stock": (
                    MISSING_SHELF_LIFE_LABEL
                    if category in PURCHASED_CATEGORIES and stock_unavailable
                    else MISSING_STOCK_LABEL
                    if stock_unavailable
                    else round(stock_qty, 1)
                ),
                "net_need": round(net_need, 1),
                "production_plan": production_plan,
                "total_for_sale": round(production_plan + stock_qty, 1),
                "kratnost": kratnost_display,
            }
        )
    rows.sort(key=lambda item: (item["bakery_id"], item["category"], item["product_name"]))
    return rows


def build_excel(rows: list[dict], forecast_date: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    d = date_type.fromisoformat(forecast_date)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прогноз"
    ws["A1"] = f"Прогноз выпечки — {d.strftime('%d.%m.%Y')} ({WEEKDAY_RU[d.weekday()]})"
    ws["A1"].font = Font(bold=True, size=12)
    headers = [
        "Пекарня",
        "Категория",
        "Номенклатура",
        "Прогноз",
        "Остаток со вчерашнего дня",
        "Чистая потребность",
        "План выпуска",
        "Итого на продажу",
        "Кратность",
    ]
    widths = [35, 20, 40, 12, 24, 20, 16, 18, 24]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 30
    fills = [PatternFill("solid", fgColor="EBF3FB"), PatternFill("solid", fgColor="FFFFFF")]
    prev_bid = None
    fill_idx = 0
    for data_row in rows:
        if data_row["bakery_id"] != prev_bid:
            fill_idx = 1 - fill_idx
            prev_bid = data_row["bakery_id"]
        row_num = ws.max_row + 1
        values = [
            data_row["bakery_name"],
            data_row["category"],
            data_row["product_name"],
            data_row["forecast"],
            data_row["yesterday_stock"],
            data_row["net_need"],
            data_row["production_plan"],
            data_row["total_for_sale"],
            data_row["kratnost"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.fill = fills[fill_idx]
            cell.border = border
            cell.font = Font(size=10)
            if col_idx >= 4 and isinstance(value, (int, float)):
                cell.number_format = "#,##0.0" if col_idx in {4, 5, 6, 8} else "#,##0"
                cell.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{ws.max_row}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_to_chat(file_bytes: bytes, filename: str, forecast_date: str) -> None:
    d = date_type.fromisoformat(forecast_date)
    weekday_name = WEEKDAY_RU[d.weekday()]
    api_key = os.environ.get("VIBECODE_API_KEY") or ""
    if not api_key:
        raise RuntimeError("VIBECODE_API_KEY not set")
    b24_webhook_base = os.environ.get(B24_WEBHOOK_URL_ENV, "").rstrip("/")
    if not b24_webhook_base:
        raise RuntimeError(f"{B24_WEBHOOK_URL_ENV} not set in environment")
    ru_filename = f"Прогноз_{d.strftime('%d.%m.%Y')}_{weekday_name}.xlsx"

    step1_body = json.dumps({
        "id": PILOT_CHAT_DISK_FOLDER_ID,
        "data": {"NAME": ru_filename},
        "generateUniqueName": "Y",
    }).encode("utf-8")
    req = urllib.request.Request(
        f"{b24_webhook_base}/disk.folder.uploadfile",
        data=step1_body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        step1 = json.loads(resp.read())
    if "error" in step1:
        raise RuntimeError(f"disk.folder.uploadfile failed: {step1}")
    upload_url = step1["result"]["uploadUrl"]
    print("  [b24] uploadUrl obtained")

    boundary = f"----FormBoundary{int(_time.time())}"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{ru_filename}"\r\n'
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n"
        "\r\n"
    ).encode("utf-8") + file_bytes + f"\r\n--{boundary}--\r\n".encode("utf-8")
    req = urllib.request.Request(
        upload_url,
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=120) as resp:
        step2 = json.loads(resp.read())
    if "error" in step2:
        raise RuntimeError(f"File upload to uploadUrl failed: {step2}")
    disk_id = step2["result"]["ID"]
    print(f"  [b24] file uploaded, disk_id={disk_id}")

    msg_body = json.dumps({"message": f"Прогноз — {d.strftime('%d.%m.%Y')} ({weekday_name})"}).encode("utf-8")
    req = urllib.request.Request(
        f"{VIBECODE_API_BASE}/chats/{PILOT_CHAT_DIALOG_ID}/messages",
        data=msg_body,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        msg_result = json.loads(resp.read())
    if not msg_result.get("success"):
        raise RuntimeError(f"Message send failed: {msg_result}")
    print(f"  [vibecode] text message sent, id={msg_result['data']}")

    commit_body = json.dumps({"CHAT_ID": PILOT_CHAT_ID, "DISK_ID": disk_id}).encode("utf-8")
    req = urllib.request.Request(
        f"{b24_webhook_base}/im.disk.file.commit",
        data=commit_body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        commit_result = json.loads(resp.read())
    if "error" in commit_result:
        raise RuntimeError(f"im.disk.file.commit failed: {commit_result}")
    print(f"  [b24] file message sent, message_id={commit_result.get('result', {}).get('MESSAGE_ID')}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--env-file", default="/opt/app/.env")
    parser.add_argument("--date", default=None, help="Forecast date YYYY-MM-DD; default: today")
    parser.add_argument("--dry-run", action="store_true", help="Build the workbook but do not send to Bitrix24")
    parser.add_argument("--out-dir", default="/opt/output/weighted_weekday_forecast")
    args = parser.parse_args()

    load_env(args.env_file)
    forecast_date = args.date or str(date_type.today())
    d = date_type.fromisoformat(forecast_date)
    weekday_abbr = WEEKDAY_RU[d.weekday()][:2]
    print(f"Weighted weekday pilot forecast | date: {forecast_date} ({weekday_abbr})")

    client = get_client()
    pilot_bakery_ids = load_pilot_bakery_ids(client)
    run_id, scope = fetch_active_scope(client, forecast_date, pilot_bakery_ids)
    history_dates = [(d - timedelta(days=7 * weeks)).isoformat() for weeks in range(1, 9)]
    history = fetch_history_demand(client, history_dates, pilot_bakery_ids)
    forecast = build_weighted_weekday(scope, history, forecast_date)
    rows = build_rows(client, forecast_date, forecast, pilot_bakery_ids)
    if not rows:
        raise RuntimeError("No publisher rows generated")

    print(f"  active scope run: {run_id}")
    print(f"  override rows: {len(forecast)}, publisher rows: {len(rows)}")
    print(f"  forecast total: {sum(float(row['forecast']) for row in rows):.1f}")
    print(f"  production plan total: {sum(float(row['production_plan']) for row in rows):.1f}")

    file_bytes = build_excel(rows, forecast_date)
    filename = f"Прогноз_выпечки_{d.strftime('%d.%m.%Y')}_{weekday_abbr}_weighted_weekday.xlsx"
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / filename
    out_path.write_bytes(file_bytes)
    print(f"  saved: {out_path}")

    if args.dry_run:
        print("  --dry-run: skipping Bitrix24 send")
        return
    send_to_chat(file_bytes, filename, forecast_date)
    print("  done.")


if __name__ == "__main__":
    main()
